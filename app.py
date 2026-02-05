import streamlit as st
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import random
import io
import re
from PIL import Image

# --- הגדרות עמוד ---
st.set_page_config(page_title="משחק הדגים", page_icon="🐠", layout="wide")

# --- עיצוב CSS ---
st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Rubik:wght@400;700&display=swap');

    .stApp {
        direction: rtl;
        background-color: #e6f2ff;
        font-family: 'Rubik', sans-serif;
    }

    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}

    /* סטטוס בר */
    .status-bar {
        background-color: #3d5afe;
        color: white;
        padding: 15px;
        border-radius: 15px;
        box-shadow: 0 4px 6px rgba(0,0,0,0.1);
        display: flex;
        justify-content: space-around;
        align-items: center;
        font-size: 22px;
        font-weight: bold;
        margin-bottom: 20px;
    }

    /* כרטיס משחק */
    .game-card {
        background-color: white;
        padding: 20px;
        border-radius: 20px;
        box-shadow: 0 8px 16px rgba(0,0,0,0.1);
        text-align: center;
        margin-bottom: 20px;
    }

    .question-text {
        font-size: 28px;
        font-weight: bold;
        color: #1565c0;
        margin: 15px 0;
    }

    /* כפתורים */
    .stButton>button {
        width: 100%;
        height: 70px;
        font-size: 24px !important;
        font-weight: bold;
        border-radius: 12px;
        background-color: #f8f9fa;
        border: 2px solid #e9ecef;
        color: #333;
        transition: all 0.2s;
    }
    .stButton>button:hover {
        background-color: #e3f2fd;
        border-color: #3d5afe;
        color: #3d5afe;
        transform: translateY(-2px);
    }

    div[data-baseweb="select"] > div {
        background-color: #e3f2fd;
        border-radius: 10px;
        border: 2px solid #bbdefb;
    }

    .primary-btn button {
        background: linear-gradient(45deg, #00c853, #69f0ae) !important;
        color: white !important;
        border: none !important;
        box-shadow: 0 4px 10px rgba(0, 200, 83, 0.4);
    }

    section[data-testid="stSidebar"] {
        background-color: #f1f8e9;
        direction: rtl;
    }

    .aquarium-title {
        text-align: center;
        color: #2e7d32;
        border-bottom: 2px solid #a5d6a7;
        padding-bottom: 10px;
        margin-bottom: 15px;
    }
    </style>
""", unsafe_allow_html=True)


# --- פונקציות לוגיקה ---

def clean_text(text):
    return text.replace(',', '').replace('.', '').replace('_', ' ').strip()


@st.cache_data
def load_data_from_pptx(pptx_path):
    """טעינת נתונים חכמה מקובץ PowerPoint"""
    if 'fish_data_cache' not in st.session_state:
        data = []
        try:
            prs = Presentation(pptx_path)

            for slide in prs.slides:
                hebrew_name = None
                family_name = None
                images = []

                # 1. איסוף כל הטקסט בשקופית
                full_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        full_text.append(shape.text)

                # ניתוח הטקסט (החלק המתוקן)
                slide_text_content = "\n".join(full_text)
                for line in slide_text_content.split('\n'):
                    clean_line = clean_text(line)
                    if not clean_line: continue

                    # בדיקת עברית (שם הדג)
                    if any("\u0590" <= c <= "\u05EA" for c in clean_line):
                        # מנקים ממנו את האנגלית כדי לקבל שם נקי
                        clean_hebrew = re.sub(r'[a-zA-Z]', '', clean_line).replace(',', '').strip()
                        if len(clean_hebrew) > 2:
                            hebrew_name = clean_hebrew

                    # בדיקת אנגלית (משפחה) - באותה שורה או בשורה נפרדת
                    english_matches = re.findall(r'[a-zA-Z]{3,}', clean_line)  # מחפש מילים באנגלית
                    if english_matches:
                        possible_family = " ".join(english_matches)
                        # מסננים מילים טכניות שלא קשורות
                        if "source" not in possible_family.lower() and "image" not in possible_family.lower():
                            family_name = possible_family

                # 2. איסוף כל התמונות בשקופית
                for shape in slide.shapes:
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        image_bytes = shape.image.blob
                        img = Image.open(io.BytesIO(image_bytes))
                        images.append(img)

                # שמירה רק אם מצאנו גם שם וגם תמונות
                if hebrew_name and images:
                    data.append({
                        "name": hebrew_name,
                        "family": family_name if family_name else "לא ידוע",  # ברירת מחדל אם לא נמצא
                        "images": images
                    })

            return data

        except Exception as e:
            st.error(f"שגיאה בטעינת הקובץ: {e}")
            return []
    return []


def mask_fish_name(full_name):
    clean_n = full_name.replace('-', ' ')
    words = clean_n.split()
    if len(words) == 1:
        return f"{clean_n[0]}_____", clean_n[1:]

    index_to_hide = random.randint(0, len(words) - 1)
    missing_word = words[index_to_hide]
    words[index_to_hide] = "_____"
    display_text = " ".join(words)
    return display_text, missing_word


# --- ניהול משחק ---

def init_game():
    if 'game_started' not in st.session_state:
        st.session_state.game_started = False
    if 'aquarium' not in st.session_state:
        st.session_state.aquarium = []


def start_game(selected_mode, initial_lives):
    st.session_state.game_started = True
    st.session_state.mode = selected_mode
    st.session_state.score = 0
    st.session_state.lives = initial_lives
    st.session_state.game_over = False
    st.session_state.aquarium = []
    new_turn(data)


def new_turn(data):
    fish = random.choice(data)
    st.session_state.current_fish = fish
    st.session_state.answered = False
    st.session_state.feedback = None
    st.session_state.round_score = 0
    st.session_state.round_feedback = []

    if 'turn_count' not in st.session_state:
        st.session_state.turn_count = 0
    st.session_state.turn_count += 1

    if st.session_state.mode == "multiple_choice":
        all_names = list(set([d['name'] for d in data]))
        if fish['name'] in all_names: all_names.remove(fish['name'])
        distractors = random.sample(all_names, 3)
        options = distractors + [fish['name']]
        random.shuffle(options)
        st.session_state.options = options
    else:
        display_text, missing_word = mask_fish_name(fish['name'])
        st.session_state.masked_text = display_text
        st.session_state.missing_word = missing_word


def check_answer(user_name_answer, user_family_answer):
    st.session_state.answered = True
    correct_name = False
    correct_family = False

    # בדיקת שם
    if st.session_state.mode == "multiple_choice":
        if user_name_answer == st.session_state.current_fish['name']:
            correct_name = True
    else:
        clean_user = clean_text(user_name_answer)
        clean_correct = clean_text(st.session_state.missing_word)
        if clean_user == clean_correct:
            correct_name = True

    # בדיקת משפחה
    real_family = st.session_state.current_fish['family']
    if user_family_answer == real_family:
        correct_family = True

    points = 0
    feedback_msg = []

    if correct_name:
        points += 10
        feedback_msg.append("✅ שם הדג נכון!")
    else:
        feedback_msg.append(f"❌ טעות בשם הדג... זהו {st.session_state.current_fish['name']}")

    if correct_family:
        points += 5
        feedback_msg.append("✅ משפחה נכונה!")
    else:
        feedback_msg.append(f"❌ טעות במשפחה... זוהי {real_family}")

    st.session_state.score += points
    st.session_state.round_feedback = feedback_msg

    if not correct_name:
        st.session_state.lives -= 1
        st.session_state.feedback = "wrong"
        if st.session_state.lives == 0:
            st.session_state.game_over = True
    else:
        st.session_state.feedback = "correct"
        if not any(f['name'] == st.session_state.current_fish['name'] for f in st.session_state.aquarium):
            st.session_state.aquarium.append(st.session_state.current_fish)
        if len(st.session_state.aquarium) % 5 == 0:
            st.balloons()


# --- טעינה ---
init_game()

# טוען קובץ PPTX
data = load_data_from_pptx("fish.pptx")

if not data:
    st.warning("לא נמצאו דגים. וודא שהקובץ fish.pptx קיים בתיקייה.")
    st.stop()

# יצירת רשימת משפחות נקייה מ-None או undefined
all_families = sorted(list(set([d['family'] for d in data if d['family'] and d['family'] != "לא ידוע"])))

# --- סרגל צד ---
with st.sidebar:
    st.markdown('<h2 class="aquarium-title">🐠 האקווריום שלי</h2>', unsafe_allow_html=True)
    if st.session_state.aquarium:
        st.write(f"**סה\"כ נאספו: {len(st.session_state.aquarium)}**")
        cols = st.columns(2)
        for i, fish in enumerate(st.session_state.aquarium):
            with cols[i % 2]:
                st.image(fish['images'][0], use_container_width=True)
                st.caption(f"{fish['name']}")
    else:
        st.info("האקווריום ריק... תתחילו לשחק כדי לאסוף דגים!")

# --- תוכן ראשי ---
main_col1, main_content, main_col3 = st.columns([1, 2, 1])

with main_content:
    if not st.session_state.game_started:
        st.markdown("""
            <div class="game-card">
                <h1>🌊 ברוכים הבאים למשחק הדגים!</h1>
                <p>עכשיו עם תמונות כפולות לזיהוי מדויק יותר!</p>
            </div>
        """, unsafe_allow_html=True)

        col_sets1, col_sets2 = st.columns(2)
        with col_sets1:
            st.markdown("### 🎮 סוג משחק")
            mode = st.radio("", ("שאלות אמריקאיות", "השלם את החסר"), label_visibility="collapsed")
        with col_sets2:
            st.markdown("### ❤️ כמות חיים")
            lives_val = st.slider("", 1, 10, 3, label_visibility="collapsed")

        selected_mode_key = "multiple_choice" if "אמריקאי" in mode else "text_input"
        st.write("---")

        st.markdown('<div class="primary-btn">', unsafe_allow_html=True)
        if st.button("התחל משחק 🚀", use_container_width=True):
            start_game(selected_mode_key, lives_val)
            st.rerun()
        st.markdown('</div>', unsafe_allow_html=True)

    elif st.session_state.game_over:
        st.markdown(f"""
            <div class="game-card">
                <h1 style="color: #d32f2f;">😢 המשחק נגמר!</h1>
                <h2>ניקוד סופי: {st.session_state.score}</h2>
            </div>
        """, unsafe_allow_html=True)

        st.markdown('<div class="primary-btn">', unsafe_allow_html=True)
        if st.button("חזרה למסך הראשי 🏠", use_container_width=True):
            st.session_state.game_started = False
            st.rerun()
        st.markdown('</div>', unsafe_allow_html=True)

    else:
        hearts = "❤️" * st.session_state.lives
        st.markdown(f"""
            <div class="status-bar">
                <div>⭐ {st.session_state.score}</div>
                <div>{hearts}</div>
            </div>
        """, unsafe_allow_html=True)

        st.markdown('<div class="game-card">', unsafe_allow_html=True)

        # תצוגת תמונות
        current_images = st.session_state.current_fish['images']
        num_images = len(current_images)

        if num_images > 0:
            cols = st.columns(num_images)
            for idx, img in enumerate(current_images):
                with cols[idx]:
                    st.image(img, use_container_width=True)

        if st.session_state.answered:
            for msg in st.session_state.round_feedback:
                if "✅" in msg:
                    st.success(msg)
                else:
                    st.error(msg)

            st.markdown('<div class="primary-btn">', unsafe_allow_html=True)
            if st.button("לדג הבא ➡️", use_container_width=True):
                new_turn(data)
                st.rerun()
            st.markdown('</div>', unsafe_allow_html=True)

        else:
            st.markdown("### 1. לאיזה משפחה הדג שייך?")
            family_key = f"fam_{st.session_state.turn_count}"
            selected_family = st.selectbox("בחרי משפחה מהרשימה:", ["בחר..."] + all_families, key=family_key,
                                           label_visibility="collapsed")

            st.markdown("---")

            if st.session_state.mode == "multiple_choice":
                st.markdown("### 2. מה שם הדג?")
                cols = st.columns(2)
                for i, opt in enumerate(st.session_state.options):
                    if cols[i % 2].button(opt, key=f"btn_{i}", use_container_width=True):
                        if selected_family == "בחר...":
                            st.toast("⚠️ נא לבחור משפחה לפני שעונים!", icon="⚠️")
                        else:
                            check_answer(opt, selected_family)
                            st.rerun()

            else:
                st.markdown(f'<div class="question-text">2. {st.session_state.masked_text}</div>',
                            unsafe_allow_html=True)
                with st.form(key='answer_form'):
                    user_input = st.text_input("השלימי את המילה החסרה:", key="user_input_field")
                    submit_button = st.form_submit_button(label='בדיקה ✅')

                    if submit_button:
                        if selected_family == "בחר...":
                            st.warning("⚠️ נא לבחור משפחה מהרשימה למעלה!")
                        elif not user_input:
                            st.warning("⚠️ נא לכתוב את שם הדג!")
                        else:
                            check_answer(user_input, selected_family)
                            st.rerun()

        st.markdown('</div>', unsafe_allow_html=True)